# -*- coding: utf-8 -*-
"""生成《作品提交表》：表格正文 2 页 + 5 页作品截图（一页一张，去多余分页）。"""
import sys, os, docx
sys.stdout.reconfigure(encoding='utf-8')
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from PIL import Image

ROOT = r"F:\zcode\suiyi\ggg"
TPL  = os.path.join(ROOT, r"site-dump\assets\files\uploads\file\20260924\20260924102044c59e90534.docx")
SRC  = os.path.join(ROOT, r"参赛方案\02-提交材料清单.md")
MAT  = os.path.join(ROOT, r"参赛方案\提交材料")
ATT  = os.path.join(ROOT, "过程截图", "2026-10-15-v7e", "提交表配图")
OUTD = os.path.join(MAT, r"作品提交表-职引.docx")
SITE = "https://aiagent.cqut.edu.cn:31134/chat/HeYM5jZOBsKPURdj"
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
PH = "【待填】"
IMGS = [
    ("配图-01-首页.png", "图 1　首页与开场白：简历诊断 · 岗位匹配 · 模拟面试 · 政策咨询"),
    ("配图-02-简历诊断.png", "图 2　简历诊断报告：分项得分、加权算式、综合匹配度与风险点复核"),
    ("配图-03-岗位匹配.png", "图 3　岗位匹配建议：匹配岗位表格、优先行动与知识库依据"),
    ("配图-04-面试复盘.png", "图 4　模拟面试复盘：四维评分、引用原话与加权算式"),
    ("配图-05-政策咨询.png", "图 5　政策咨询：可申请补贴清单与重庆市官方文号"),
]

md = open(SRC, encoding="utf-8").read().replace("\r\n", "\n")
lines = md.split("\n")
def grab(head):
    i = next(i for i, l in enumerate(lines) if l.startswith(head))
    out = []
    for l in lines[i+1:]:
        if l.startswith("###") or l.startswith("---"): break
        out.append(l)
    return out
def strip(ls):
    ls = [l for l in ls if not l.strip().startswith(">")]
    while ls and not ls[0].strip(): ls.pop(0)
    while ls and not ls[-1].strip(): ls.pop()
    return ls
def curly(s):
    o, op = [], True
    for ch in s:
        if ch == '"': o.append("\u201c" if op else "\u201d"); op = not op
        else: o.append(ch)
    return "".join(o)
A = curly(strip(grab("### 草稿 A"))[0])
B = [curly(x) for x in strip(grab("### 草稿 B"))]
C = [curly(x) for x in strip(grab("### 草稿 C"))]

def set_tc(tc, paragraphs, align=WD_ALIGN_PARAGRAPH.CENTER, size=10.5,
           font="方正仿宋_GBK", line_spacing=None):
    from docx.table import _Cell
    cell = _Cell(tc, None)
    for p in cell.paragraphs[1:]:
        p._element.getparent().remove(p._element)
    p0 = cell.paragraphs[0]
    for r in list(p0.runs): r._element.getparent().remove(r._element)
    for i, txt in enumerate(paragraphs):
        p = p0 if i == 0 else cell.add_paragraph()
        p.alignment = align
        if line_spacing is not None: p.paragraph_format.line_spacing = line_spacing
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(txt); r.font.size = Pt(size); r.font.name = font
        r._element.rPr.rFonts.set(W+"eastAsia", font)
        r._element.rPr.rFonts.set(W+"ascii", font)
        r._element.rPr.rFonts.set(W+"hAnsi", font)

d = docx.Document(TPL)
tb = d.tables[0]
def T(ri, ti): return tb.rows[ri]._tr.findall(W+"tc")[ti]

set_tc(T(1, 1), ["职引小队"])
set_tc(T(2, 1), [PH]); set_tc(T(2, 3), [PH]); set_tc(T(2, 5), [PH])
set_tc(T(3, 1), [PH + "（单人参赛，与队长一致）"])
set_tc(T(4, 1), ["无（选填）"])
set_tc(T(6, 1), ["\u201c职引\u201d\u2014\u2014面向高校学生的就业智能体"])
set_tc(T(6, 3), ["管理服务智能体"])
set_tc(T(7, 1), [A], align=WD_ALIGN_PARAGRAPH.LEFT, line_spacing=1.15)
set_tc(T(8, 1), B, align=WD_ALIGN_PARAGRAPH.LEFT, line_spacing=1.0)
set_tc(T(9, 1), C, align=WD_ALIGN_PARAGRAPH.LEFT, line_spacing=1.0)
set_tc(T(10, 1), [SITE], size=11)

sec = d.sections[0]
usable_w_cm = (sec.page_width - sec.left_margin - sec.right_margin) / 360000
usable_h_cm = (sec.page_height - sec.top_margin - sec.bottom_margin) / 360000

# 首页只放「附件」标题 + 第一张图；后续图各自跟随分页
tb_head = d.add_paragraph(); tb_head.alignment = WD_ALIGN_PARAGRAPH.CENTER
tr = tb_head.add_run("附件：作品截图（5 张）")
tr.font.size = Pt(13); tr.font.bold = True; tr.font.name = "仿宋"
tr._element.rPr.rFonts.set(W+"eastAsia", "仿宋")
tb_head.paragraph_format.keep_with_next = True
tb_head.paragraph_format.space_after = Pt(6)

IMG_H_TARGET = usable_h_cm - 2.4   # 标题（首页）留 2.4cm；后续页只留标题 1.2cm 也够
for i, (fn, cap) in enumerate(IMGS, 1):
    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if i > 1:
        p.paragraph_format.page_break_before = True
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(cap); r.font.size = Pt(11); r.font.name = "仿宋"
    r._element.rPr.rFonts.set(W+"eastAsia", "仿宋")
    pic = d.add_paragraph(); pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pic.paragraph_format.space_after = Pt(0)
    iw, ih = Image.open(os.path.join(ATT, fn)).size
    ratio = ih / iw
    h_cap = IMG_H_TARGET if i == 1 else (usable_h_cm - 1.2)
    w_cm = min(usable_w_cm, h_cap / ratio)
    pic.add_run().add_picture(os.path.join(ATT, fn), width=Cm(w_cm))
    print(f"  图{i} {fn} -> 宽 {w_cm:.2f}cm 高 {w_cm*ratio:.2f}cm")

d.save(OUTD)
print("docx:", os.path.getsize(OUTD), "bytes | images:", len(docx.Document(OUTD).inline_shapes))
